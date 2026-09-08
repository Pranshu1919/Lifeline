import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def build_abstract_docx():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run('LifeLine: Smart Disaster Response and Resource Allocation System')
    run.font.name = 'Calibri'
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 73, 125)

    # Authors
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run('Pranshu, Priyanshu Kushwaha, Vijay Katiyar, and Ankita Singh')
    run.font.bold = True
    run.font.size = Pt(11)
    run.font.name = 'Calibri'

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run('Department of Information Technology, Ajay Kumar Garg Engineering College, Ghaziabad, India')
    run.font.italic = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(100, 100, 100)

    # Abstract Heading
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run('Abstract')
    run.font.name = 'Calibri'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 73, 125)

    # Abstract Paragraphs
    abstract_paragraphs = [
        "Natural disasters such as floods, earthquakes, landslides, and fires cause sudden loss of life and heavy damage to homes and roads. During a disaster, the biggest problem is getting accurate information quickly and sending help to the right places. Emergency hotlines often get overwhelmed with thousands of phone calls, leading to phone line congestion, duplicate reports, and delayed rescue operations.",
        "To solve these real-world problems, we built LifeLine—a smart, easy-to-use digital system designed to speed up disaster response and save lives. LifeLine allows affected citizens to report emergencies directly through a simple web interface by submitting text descriptions, photos, and location coordinates.",
        "The system uses Machine Learning (ML) to analyze incoming reports automatically. It identifies the disaster type (such as flood or fire), checks photos to confirm real damage, filters out fake or duplicate reports, and assigns an emergency urgency rating from Level 1 to Level 5. All verified emergencies are shown on a live interactive map, giving emergency commanders a clear view of affected areas, hospitals, and available rescue units.",
        "Finally, LifeLine uses a smart resource allocation algorithm to calculate the best routes and automatically assign the nearest available rescue teams, ambulances, and fire trucks to critical incidents. In real-world dataset tests, LifeLine identified emergency priority levels with 93.4% accuracy and reduced rescue response times by up to 38.6% compared to traditional dispatch methods."
    ]

    for para_text in abstract_paragraphs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(para_text)
        run.font.name = 'Calibri'
        run.font.size = Pt(10.5)

    # Keywords
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(12)
    run_kw = p.add_run('Keywords— ')
    run_kw.font.bold = True
    run_kw.font.name = 'Calibri'
    run_kw.font.size = Pt(10)
    
    run_val = p.add_run('Disaster Management, Emergency Response, Machine Learning, Resource Allocation, Smart Rescue Dispatch, Interactive Maps.')
    run_val.font.italic = True
    run_val.font.name = 'Calibri'
    run_val.font.size = Pt(10)

    try:
        doc.save('LifeLine_Abstract_SimpleEnglish.docx')
        print('Saved LifeLine_Abstract_SimpleEnglish.docx!')
    except Exception as e:
        print('Error saving:', e)

if __name__ == '__main__':
    build_abstract_docx()
